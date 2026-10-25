import io
import re

from docx import Document as DocxDocument
from pypdf import PdfReader


def extract_text(filename, raw_bytes):
    lower = filename.lower()
    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if lower.endswith(".docx"):
        doc = DocxDocument(io.BytesIO(raw_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    return raw_bytes.decode("utf-8")


def chunk_document(text, max_chars=600, overlap=100):
    sentences = re.split(r"(?<=[.!?])\s+", text.strip())
    chunks = []
    current = ""
    for sentence in sentences:
        if not sentence.strip():
            continue
        candidate = (current + " " + sentence).strip() if current else sentence.strip()
        if len(candidate) > max_chars and current:
            chunks.append(current)
            current = current[-overlap:] + " " + sentence.strip() if len(current) > overlap else sentence.strip()
            current = current.strip()
        else:
            current = candidate
    if current:
        chunks.append(current)
    return [{"index": i, "text": chunk} for i, chunk in enumerate(chunks)]
